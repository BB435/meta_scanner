"""Extract native text and embedded image text without office conversion."""

from __future__ import annotations

import re
import sqlite3
import zipfile
from collections.abc import Callable
from dataclasses import asdict, dataclass
from pathlib import Path
from xml.etree import ElementTree

from .config import Settings
from .ocr import OcrError, recognize_image
from .repository import OcrCheckpoint


class DocumentError(RuntimeError):
    def __init__(self, code: str, detail: str):
        self.code = code
        super().__init__(detail)


class ProcessingDeferred(RuntimeError):
    """The processing window ended between extraction units."""


@dataclass(frozen=True)
class TextBlock:
    locator: str
    text: str
    method: str

    def to_dict(self) -> dict:
        return asdict(self)


@dataclass
class Extraction:
    blocks: list[TextBlock]
    warnings: list[str]


class TextBlocks(list[TextBlock]):
    def __init__(self, max_chars: int):
        super().__init__()
        self.max_chars = max_chars
        self.total_chars = 0

    def append(self, block: TextBlock) -> None:
        self.total_chars += len(block.text)
        if self.total_chars > self.max_chars:
            raise DocumentError("LIMIT_EXCEEDED", "Extracted text exceeds limit")
        super().append(block)


def _text(value: object) -> str:
    return re.sub(r"[\t \u3000]+", " ", "" if value is None else str(value)).strip()


def _append(
    blocks: list[TextBlock],
    locator: str,
    value: object,
    method: str,
    settings: Settings,
) -> None:
    text = _text(value)
    if not text:
        return
    blocks.append(TextBlock(locator, text, method))


def _office_zip(path: Path, extension: str, settings: Settings) -> zipfile.ZipFile:
    if not zipfile.is_zipfile(path):
        with path.open("rb") as source:
            if source.read(8) == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
                raise DocumentError(
                    "PROTECTED", "Encrypted or unsupported Office container"
                )
        raise DocumentError("UNSUPPORTED_CONTENT", "Office document is not OOXML")
    archive = zipfile.ZipFile(path)
    expected = {
        ".docx": "word/document.xml",
        ".pptx": "ppt/presentation.xml",
        ".xlsx": "xl/workbook.xml",
    }[extension]
    names = set(archive.namelist())
    if expected not in names:
        archive.close()
        raise DocumentError(
            "UNSUPPORTED_CONTENT", "File content does not match extension"
        )
    total = sum(entry.file_size for entry in archive.infolist())
    if total > settings.max_uncompressed_bytes:
        archive.close()
        raise DocumentError("LIMIT_EXCEEDED", "OOXML expanded size exceeds limit")
    return archive


def _has_element(archive: zipfile.ZipFile, name: str, element: str) -> bool:
    try:
        with archive.open(name) as source:
            for event, node in ElementTree.iterparse(source, events=("start", "end")):
                if node.tag.rsplit("}", 1)[-1] == element:
                    return True
                if event == "end":
                    node.clear()
    except KeyError:
        return False
    except ElementTree.ParseError as exc:
        raise DocumentError("CORRUPT", f"Invalid XML in {name}: {exc}") from exc
    return False


def _has_enabled_protection(
    archive: zipfile.ZipFile,
    name: str,
    element: str,
    enabled_attributes: tuple[str, ...],
) -> bool:
    try:
        with archive.open(name) as source:
            for _, node in ElementTree.iterparse(source, events=("start",)):
                if node.tag.rsplit("}", 1)[-1] != element:
                    continue
                return any(
                    node.attrib.get(attribute, "").casefold() in {"1", "true"}
                    for attribute in enabled_attributes
                )
    except KeyError:
        return False
    except ElementTree.ParseError as exc:
        raise DocumentError("CORRUPT", f"Invalid XML in {name}: {exc}") from exc
    return False


def _check_protection(archive: zipfile.ZipFile, extension: str) -> None:
    if extension == ".docx" and _has_element(
        archive, "word/settings.xml", "documentProtection"
    ):
        raise DocumentError("PROTECTED", "Word editing protection is enabled")
    if extension == ".xlsx":
        if _has_enabled_protection(
            archive,
            "xl/workbook.xml",
            "workbookProtection",
            ("lockStructure", "lockWindows", "lockRevision"),
        ):
            raise DocumentError("PROTECTED", "Workbook protection is enabled")
        for name in archive.namelist():
            protected_sheet = re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name)
            if protected_sheet and _has_enabled_protection(
                archive,
                name,
                "sheetProtection",
                ("sheet", "objects", "scenarios"),
            ):
                raise DocumentError("PROTECTED", f"Sheet protection is enabled: {name}")
    if extension == ".pptx" and _has_element(
        archive, "ppt/presentation.xml", "modifyVerifier"
    ):
        raise DocumentError("PROTECTED", "PowerPoint editing protection is enabled")


def _embedded_images(
    archive: zipfile.ZipFile,
    extension: str,
    settings: Settings,
    blocks: list[TextBlock],
    warnings: list[str],
    checkpoint: OcrCheckpoint | None,
    check_deadline: Callable[[], None] | None,
) -> None:
    if not settings.ocr_enabled:
        return
    prefix = {".docx": "word/media/", ".pptx": "ppt/media/", ".xlsx": "xl/media/"}[
        extension
    ]
    for item in archive.infolist():
        if check_deadline:
            check_deadline()
        if not item.filename.startswith(prefix) or item.is_dir():
            continue
        if item.file_size > settings.max_file_bytes:
            warnings.append(f"LIMIT_EXCEEDED:{item.filename}")
            continue
        locator = f"media:{item.filename}"
        recognized = checkpoint.get(locator) if checkpoint else None
        try:
            if recognized is None:
                recognized = recognize_image(
                    archive.read(item), settings, check_deadline=check_deadline
                )
                if checkpoint:
                    checkpoint.save(locator, recognized)
        except (OcrError, OSError, ValueError) as exc:
            warnings.append(f"OCR_FAILED:{item.filename}:{exc}")
            continue
        _append(blocks, locator, recognized, "ocr", settings)


def _extract_docx(path: Path, settings: Settings, blocks: list[TextBlock]) -> None:
    from docx import Document
    from docx.table import Table

    def append_tables(tables, locator_prefix: str) -> None:
        for table in tables:
            for row_index, row in enumerate(table.rows, 1):
                values = [cell.text.strip() for cell in row.cells]
                _append(
                    blocks,
                    f"{locator_prefix}/row:{row_index}",
                    " | ".join(values),
                    "native",
                    settings,
                )

    document = Document(str(path.absolute()))
    for index, item in enumerate(document.iter_inner_content(), 1):
        if isinstance(item, Table):
            append_tables([item], f"table:{index}")
        else:
            _append(blocks, f"paragraph:{index}", item.text, "native", settings)
    for section_index, section in enumerate(document.sections, 1):
        for kind in ("header", "footer"):
            story = getattr(section, kind)
            for paragraph_index, paragraph in enumerate(story.paragraphs, 1):
                _append(
                    blocks,
                    f"section:{section_index}/{kind}:{paragraph_index}",
                    paragraph.text,
                    "native",
                    settings,
                )
            for table_index, table in enumerate(story.tables, 1):
                append_tables(
                    [table], f"section:{section_index}/{kind}/table:{table_index}"
                )


def _walk_shapes(
    shapes,
    slide_number: int,
    settings: Settings,
    blocks: list[TextBlock],
    prefix: str = "",
) -> None:
    for index, shape in enumerate(shapes, 1):
        locator = f"slide:{slide_number}/shape:{prefix}{index}"
        if hasattr(shape, "text"):
            _append(blocks, locator, shape.text, "native", settings)
        if getattr(shape, "has_table", False):
            for row_index, row in enumerate(shape.table.rows, 1):
                _append(
                    blocks,
                    f"{locator}/row:{row_index}",
                    " | ".join(cell.text for cell in row.cells),
                    "native",
                    settings,
                )
        if hasattr(shape, "shapes"):
            _walk_shapes(
                shape.shapes, slide_number, settings, blocks, prefix=f"{prefix}{index}."
            )


def _extract_pptx(
    path: Path, settings: Settings, blocks: list[TextBlock], warnings: list[str]
) -> None:
    from pptx import Presentation

    presentation = Presentation(str(path.absolute()))
    for number, slide in enumerate(presentation.slides, 1):
        if number > settings.max_pages:
            raise DocumentError("LIMIT_EXCEEDED", "Slide count exceeds limit")
        _walk_shapes(slide.shapes, number, settings, blocks)
        if settings.include_notes and slide.has_notes_slide:
            try:
                _append(
                    blocks,
                    f"slide:{number}/notes",
                    slide.notes_slide.notes_text_frame.text,  # pyright: ignore[reportOptionalMemberAccess]
                    "native",
                    settings,
                )
            except AttributeError:
                warnings.append(f"NOTES_UNAVAILABLE:slide:{number}")


def _extract_xlsx(
    path: Path, settings: Settings, blocks: list[TextBlock], warnings: list[str]
) -> None:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    formula_workbook = load_workbook(
        path, read_only=True, data_only=False, keep_links=False
    )
    value_workbook = load_workbook(
        path, read_only=True, data_only=True, keep_links=False
    )
    try:
        count = 0
        has_formulas = False
        value_sheets = {sheet.title: sheet for sheet in value_workbook}
        for sheet in formula_workbook:
            if sheet.sheet_state != "visible" and not settings.include_hidden_sheets:
                continue
            value_sheet = value_sheets[sheet.title]
            value_rows = value_sheet.iter_rows()
            for row_index, row in enumerate(sheet.iter_rows(), 1):
                cached_row = next(value_rows)
                pairs = []
                for column_index, cell in enumerate(row, 1):
                    cached_cell = cached_row[column_index - 1]
                    if cell.data_type == "f":
                        has_formulas = True
                        if cached_cell.value is None:
                            warnings.append(
                                f"FORMULA_CACHE_MISSING:sheet:{sheet.title}/{cell.coordinate}"
                            )
                            value = f"{cell.value} (計算結果なし)"
                        else:
                            value = cached_cell.value
                    else:
                        value = cached_cell.value
                    if value is None or value == "":
                        continue
                    count += 1
                    if count > settings.max_cells:
                        raise DocumentError(
                            "LIMIT_EXCEEDED", "Nonempty cell count exceeds limit"
                        )
                    pairs.append(
                        f"{get_column_letter(column_index)}{row_index}:{value}"
                    )
                if pairs:
                    _append(
                        blocks,
                        f"sheet:{sheet.title}/row:{row_index}",
                        " | ".join(pairs),
                        "native",
                        settings,
                    )
        if has_formulas:
            warnings.append("FORMULA_CACHE_ONLY:formula results are not recalculated")
    finally:
        formula_workbook.close()
        value_workbook.close()


def _extract_pdf(
    path: Path,
    settings: Settings,
    blocks: list[TextBlock],
    warnings: list[str],
    checkpoint: OcrCheckpoint | None,
    check_deadline: Callable[[], None] | None,
) -> None:
    from pypdf import PdfReader

    reader = PdfReader(path, strict=False)
    if reader.is_encrypted:
        raise DocumentError("PROTECTED", "Encrypted PDF")
    if len(reader.pages) > settings.max_pages:
        raise DocumentError("LIMIT_EXCEEDED", "PDF page count exceeds limit")
    pdfium_document = None
    if settings.ocr_enabled:
        import pypdfium2 as pdfium

        pdfium_document = pdfium.PdfDocument(str(path))
    try:
        for number, page in enumerate(reader.pages, 1):
            if check_deadline:
                check_deadline()
            try:
                native = page.extract_text() or ""
            except Exception as exc:  # noqa: BLE001
                native = ""
                warnings.append(f"TEXT_LAYER_FAILED:page:{number}:{type(exc).__name__}")
            _append(blocks, f"page:{number}", native, "native", settings)
            if pdfium_document is None:
                continue
            try:
                locator = f"page:{number}/ocr"
                recognized = checkpoint.get(locator) if checkpoint else None
                if recognized is None:
                    pdfium_page = pdfium_document[number - 1]
                    try:
                        scale = settings.dpi / 72
                        if (
                            pdfium_page.get_width()
                            * pdfium_page.get_height()
                            * scale
                            * scale
                            > settings.max_page_pixels
                        ):
                            warnings.append(
                                f"LIMIT_EXCEEDED:page:{number}:OCR pixel limit"
                            )
                            continue
                        rendered = pdfium_page.render(scale=scale)  # pyright: ignore[reportArgumentType]
                        try:
                            recognized = recognize_image(
                                rendered.to_pil(),
                                settings,
                                check_deadline=check_deadline,
                            )
                        finally:
                            rendered.close()
                    finally:
                        pdfium_page.close()
                    if checkpoint:
                        checkpoint.save(locator, recognized)
                native_lines = {
                    re.sub(r"\s+", "", line)
                    for line in native.splitlines()
                    if line.strip()
                }
                fresh = [
                    line
                    for line in recognized.splitlines()
                    if line.strip() and re.sub(r"\s+", "", line) not in native_lines
                ]
                _append(blocks, locator, "\n".join(fresh), "ocr", settings)
            except DocumentError, ProcessingDeferred, sqlite3.Error:
                raise
            except Exception as exc:  # noqa: BLE001
                warnings.append(f"OCR_FAILED:page:{number}:{exc}")
    finally:
        if pdfium_document is not None:
            pdfium_document.close()


def _extract_text(path: Path, settings: Settings, blocks: list[TextBlock]) -> None:
    data = path.read_bytes()
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        if not settings.allow_cp932:
            raise DocumentError("DECODE_ERROR", "Text is not UTF-8") from exc
        try:
            text = data.decode("cp932")
        except UnicodeDecodeError as fallback_exc:
            raise DocumentError(
                "DECODE_ERROR", "Text cannot be decoded as CP932"
            ) from fallback_exc
    for number, line in enumerate(text.splitlines(), 1):
        _append(blocks, f"line:{number}", line, "native", settings)


def extract_document(
    path: Path,
    settings: Settings,
    *,
    checkpoint: OcrCheckpoint | None = None,
    check_deadline: Callable[[], None] | None = None,
) -> Extraction:
    extension = path.suffix.lower()
    blocks: list[TextBlock] = TextBlocks(settings.max_chars)
    warnings: list[str] = []
    if extension in {".docx", ".pptx", ".xlsx"}:
        try:
            with _office_zip(path, extension, settings) as archive:
                _check_protection(archive, extension)
                if extension == ".docx":
                    _extract_docx(path, settings, blocks)
                elif extension == ".pptx":
                    _extract_pptx(path, settings, blocks, warnings)
                else:
                    _extract_xlsx(path, settings, blocks, warnings)
                _embedded_images(
                    archive,
                    extension,
                    settings,
                    blocks,
                    warnings,
                    checkpoint,
                    check_deadline,
                )
        except ProcessingDeferred, sqlite3.Error:
            raise
        except DocumentError:
            raise
        except Exception as exc:
            raise DocumentError(
                "CORRUPT", f"Cannot read Office document: {exc}"
            ) from exc
    elif extension == ".pdf":
        try:
            _extract_pdf(path, settings, blocks, warnings, checkpoint, check_deadline)
        except ProcessingDeferred, sqlite3.Error:
            raise
        except DocumentError:
            raise
        except Exception as exc:
            raise DocumentError("CORRUPT", f"Cannot read PDF: {exc}") from exc
    elif extension in {".txt", ".md"}:
        _extract_text(path, settings, blocks)
    else:
        raise DocumentError(
            "UNSUPPORTED_CONTENT", f"Unsupported extension: {extension}"
        )
    if not blocks:
        warnings.append("NO_TEXT")
    return Extraction(blocks, list(dict.fromkeys(warnings)))
