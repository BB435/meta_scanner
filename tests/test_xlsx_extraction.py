from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from meta_scanner.extractors import DocumentError, extract_document


def _settings() -> SimpleNamespace:
    return SimpleNamespace(
        max_uncompressed_bytes=10 * 1024 * 1024,
        max_chars=100_000,
        max_cells=10_000,
        include_hidden_sheets=True,
        ocr_enabled=False,
    )


def test_preserves_formula_when_cached_result_is_missing(tmp_path):
    path = tmp_path / "missing-cache.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "売上"
    sheet["A1"] = "合計"
    sheet["B1"] = "=SUM(B2:B5)"
    workbook.save(path)

    extraction = extract_document(path, _settings())

    row = next(
        block for block in extraction.blocks if block.locator == "sheet:売上/row:1"
    )
    assert row.text == "A1:合計 | B1:=SUM(B2:B5) (計算結果なし)"
    assert "FORMULA_CACHE_MISSING:sheet:売上/B1" in extraction.warnings
    assert (
        "FORMULA_CACHE_ONLY:formula results are not recalculated" in extraction.warnings
    )


def test_does_not_emit_formula_cache_warning_without_formulas(tmp_path):
    path = tmp_path / "values-only.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "値"
    workbook.save(path)

    extraction = extract_document(path, _settings())

    assert not any(
        warning.startswith("FORMULA_CACHE") for warning in extraction.warnings
    )


def test_skips_workbook_with_enabled_sheet_protection(tmp_path):
    path = tmp_path / "protected.xlsx"
    workbook = Workbook()
    workbook.active.protection.sheet = True
    workbook.save(path)

    with pytest.raises(DocumentError, match="Sheet protection is enabled") as error:
        extract_document(path, _settings())

    assert error.value.code == "PROTECTED"
