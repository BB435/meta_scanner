from types import SimpleNamespace

from docx import Document
from docx.shared import Inches

from meta_scanner.extractors import extract_document


def _settings() -> SimpleNamespace:
    return SimpleNamespace(
        max_uncompressed_bytes=10 * 1024 * 1024,
        max_chars=100_000,
        ocr_enabled=False,
    )


def test_extracts_docx_header_and_footer_tables(tmp_path):
    path = tmp_path / "header-footer.docx"
    document = Document()
    document.add_table(rows=1, cols=1).cell(0, 0).text = "本文の表"
    section = document.sections[0]
    section.header.paragraphs[0].text = "文書ヘッダー"
    section.header.add_table(rows=1, cols=2, width=Inches(6)).cell(0, 0).text = (
        "管理番号"
    )
    section.header.tables[0].cell(0, 1).text = "A-123"
    section.footer.paragraphs[0].text = "文書フッター"
    section.footer.add_table(rows=1, cols=1, width=Inches(6)).cell(0, 0).text = (
        "社外秘"
    )
    document.save(path)

    extraction = extract_document(path, _settings())

    by_locator = {block.locator: block.text for block in extraction.blocks}
    assert by_locator["table:1/row:1"] == "本文の表"
    assert by_locator["section:1/header:1"] == "文書ヘッダー"
    assert by_locator["section:1/header/table:1/row:1"] == "管理番号 | A-123"
    assert by_locator["section:1/footer:1"] == "文書フッター"
    assert by_locator["section:1/footer/table:1/row:1"] == "社外秘"

